"""Сборка отчёта .docx по образцу референса (титульный лист ЗГУ + разделы).

Использование: python build_report.py <spec.py> <out.docx>
spec.py определяет LAB (номер), TOPIC (тема — название лабораторной) и BLOCKS (список блоков):
  ("h", "Заголовок раздела")            — Heading 1, с новой страницы
  ("h2", "Подзаголовок")                — жирный абзац
  ("p", "Текст абзаца")                 — обычный абзац
  ("ul", ["пункт", ...])                — маркированный список
  ("ol", ["пункт", ...])                — нумерованный список (цифры в тексте)
  ("img", "path.png", "Подпись", w_cm)  — рисунок + «Рисунок N. Подпись»
  ("table", "Подпись", [заголовки], [[ячейки], ...])
  ("code", "текст кода")                — моноширинный блок
"""
import copy
import os
import runpy
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

HERE = os.path.dirname(os.path.abspath(__file__))
REF = os.path.join(HERE, "template.docx")

STUDENT = "Сулейманов А.Ш."
GROUP = "ИС-23"
DISCIPLINE = "Технологии программирования"
MAX_W_CM = 15.9
MAX_H_CM = 20.0


def set_run_font(run, size=14, bold=False, italic=False, mono=False):
    name = "Consolas" if mono else "Times New Roman"
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), name)


def fmt_par(par, align=WD_ALIGN_PARAGRAPH.JUSTIFY, first_line=True, line=1.5, after=0):
    pf = par.paragraph_format
    pf.alignment = align
    pf.line_spacing = line
    pf.space_after = Pt(after)
    pf.space_before = Pt(0)
    pf.first_line_indent = Cm(1.25) if first_line else Cm(0)


def set_text(par, text, **font):
    """Заменяет текст абзаца титульника, сохраняя оформление первого прогона."""
    runs = par.runs
    rpr = copy.deepcopy(runs[0]._element.rPr) if runs else None
    for r in runs:
        r._element.getparent().remove(r._element)
    run = par.add_run(text)
    if rpr is not None:
        run._element.insert(0, rpr)
    return run


def add_numpr(par, num_id):
    ppr = par._p.get_or_add_pPr()
    numpr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    nid = OxmlElement("w:numId")
    nid.set(qn("w:val"), str(num_id))
    numpr.append(ilvl)
    numpr.append(nid)
    ppr.append(numpr)


def add_rich(par, text, size=14, mono=False):
    """Поддержка **жирного** и `кода` внутри текста."""
    import re

    parts = re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            set_run_font(par.add_run(part[2:-2]), size=size, bold=True, mono=mono)
        elif part.startswith("`") and part.endswith("`"):
            set_run_font(par.add_run(part[1:-1]), size=size - 1, mono=True)
        else:
            set_run_font(par.add_run(part), size=size, mono=mono)


def set_cell_border(table):
    tbl = table._tbl
    tblpr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")
        borders.append(el)
    tblpr.append(borders)


def build(spec_path, out_path):
    spec = runpy.run_path(spec_path)
    lab = spec["LAB"]
    topic = spec["TOPIC"]
    blocks = spec["BLOCKS"]
    base_dir = os.path.dirname(os.path.abspath(spec_path))

    doc = Document(REF)
    body = doc.element.body
    paras = doc.paragraphs

    # Титульный лист: абзацы 0..24 референса (24 — разрыв страницы)
    set_text(paras[8], f"Дисциплина «{DISCIPLINE}»")
    set_text(paras[9], f"Отчет к лабораторной работе {lab}")
    set_run_font(set_text(paras[10], f"Тема: «{topic}»"))
    paras[10].paragraph_format.first_line_indent = Cm(0)
    set_text(paras[18], f"Студент: {STUDENT}")
    set_text(paras[19], f"Группа: {GROUP}")

    # paras[24] — разрыв страницы после оглавления (sdt); оглавление сохраняем,
    # Word обновит его на шаге finalize.ps1
    last_title = paras[24]._p
    removing = False
    for el in list(body):
        if removing and el.tag != qn("w:sectPr"):
            body.remove(el)
        if el is last_title:
            removing = True

    # автонумерация: {fig:имя_файла_без_расширения} в тексте заменяется номером рисунка
    figs = {}
    for b in blocks:
        if b[0] == "img":
            figs[os.path.splitext(os.path.basename(b[1]))[0]] = len(figs) + 1

    def refs(text):
        import re

        def sub(m):
            if m.group(1) not in figs:
                raise KeyError(f"нет рисунка {m.group(1)}")
            return str(figs[m.group(1)])

        return re.sub(r"\{fig:([\w.-]+)\}", sub, text)

    blocks = [
        (b[0], refs(b[1])) if b[0] in ("p", "h2") else
        (b[0], [refs(x) for x in b[1]]) if b[0] in ("ul", "ol") else
        (b[0], b[1], b[2], [[refs(str(c)) for c in row] for row in b[3]], *b[4:]) if b[0] == "table" else b
        for b in blocks
    ]

    fig = 0
    tab = 0
    tmp_files = []
    first_heading = True
    for b in blocks:
        kind = b[0]
        if kind == "h":
            par = doc.add_paragraph(style="Heading 1")
            fmt_par(par, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False, after=12)
            # разрыв страницы — свойство заголовка, а не отдельный абзац: иначе после
            # раздела, закончившегося ровно в конце страницы, появлялась пустая страница
            par.paragraph_format.page_break_before = not first_heading
            first_heading = False
            run = par.add_run(b[1])
            set_run_font(run, size=16, bold=True)
            run.font.color.rgb = RGBColor(0, 0, 0)
        elif kind == "h2":
            par = doc.add_paragraph()
            fmt_par(par, first_line=True)
            set_run_font(par.add_run(b[1]), bold=True)
        elif kind == "p":
            par = doc.add_paragraph()
            fmt_par(par)
            add_rich(par, b[1])
        elif kind in ("ul", "ol"):
            for i, item in enumerate(b[1], start=1):
                par = doc.add_paragraph(style="List Paragraph")
                fmt_par(par, first_line=False)
                if kind == "ul":
                    add_numpr(par, 1)
                    add_rich(par, item)
                else:
                    par.paragraph_format.left_indent = Cm(1.25)
                    add_rich(par, f"{i}. {item}")
        elif kind == "img":
            path = b[1] if os.path.isabs(b[1]) else os.path.join(base_dir, b[1])
            caption = b[2]
            w_cm = b[3] if len(b) > 3 else MAX_W_CM
            if w_cm == "rotate":
                # широкая диаграмма: поворот на 90° и вписывание во всю страницу
                rotated = os.path.join(os.path.dirname(os.path.abspath(out_path)), "_rot_" + os.path.basename(path))
                with Image.open(path) as im:
                    im.rotate(90, expand=True).save(rotated)
                path = rotated
                tmp_files.append(rotated)
                w_cm = MAX_W_CM
            with Image.open(path) as im:
                w, h = im.size
            max_h = 22.5 if b[3:4] == ("rotate",) else MAX_H_CM
            w_cm = min(w_cm, MAX_W_CM)
            if w_cm * h / w > max_h:
                w_cm = max_h * w / h
            par = doc.add_paragraph()
            fmt_par(par, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False, line=1.0, after=4)
            par.paragraph_format.keep_with_next = True
            par.add_run().add_picture(path, width=Cm(w_cm))
            fig += 1
            cap = doc.add_paragraph(style="Caption")
            fmt_par(cap, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False, after=10)
            run = cap.add_run(f"Рисунок {fig}. {caption}")
            set_run_font(run, size=14)
            run.font.color.rgb = RGBColor(0, 0, 0)
        elif kind == "table":
            caption, header, rows = b[1], b[2], b[3]
            widths = b[4] if len(b) > 4 else None
            tab += 1
            cap = doc.add_paragraph()
            fmt_par(cap, align=WD_ALIGN_PARAGRAPH.LEFT, first_line=False)
            cap.paragraph_format.keep_with_next = True
            set_run_font(cap.add_run(f"Таблица {tab}. {caption}"))
            table = doc.add_table(rows=1, cols=len(header))
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            set_cell_border(table)
            for i, htext in enumerate(header):
                cell = table.rows[0].cells[i]
                cell.text = ""
                p = cell.paragraphs[0]
                fmt_par(p, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False, line=1.0)
                set_run_font(p.add_run(str(htext)), size=12, bold=True)
            for row in rows:
                cells = table.add_row().cells
                for i, ctext in enumerate(row):
                    cells[i].text = ""
                    p = cells[i].paragraphs[0]
                    fmt_par(p, align=WD_ALIGN_PARAGRAPH.LEFT, first_line=False, line=1.0)
                    add_rich(p, str(ctext), size=12)
            if widths:
                table.autofit = False
                for row in table.rows:
                    for i, wcm in enumerate(widths):
                        row.cells[i].width = Cm(wcm)
            spacer = doc.add_paragraph()
            fmt_par(spacer, first_line=False)
        elif kind == "code":
            for line in b[1].rstrip("\n").split("\n"):
                par = doc.add_paragraph()
                fmt_par(par, align=WD_ALIGN_PARAGRAPH.LEFT, first_line=False, line=1.0)
                set_run_font(par.add_run(line if line else " "), size=10, mono=True)
        else:
            raise ValueError(f"unknown block {kind}")

    # убрать из пакета картинки референса, на которые больше нет ссылок
    xml = doc.element.xml
    part = doc.part
    for rid, rel in list(part.rels.items()):
        if rel.reltype.endswith("/image") and f'"{rid}"' not in xml:
            del part.rels[rid]

    doc.core_properties.author = "Сулейманов А.Ш."
    doc.core_properties.title = f"Отчет к лабораторной работе {lab}"
    os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
    doc.save(out_path)
    for f in tmp_files:
        os.remove(f)
    print("saved", out_path, "figures:", fig, "tables:", tab)


if __name__ == "__main__":
    build(sys.argv[1], sys.argv[2])
